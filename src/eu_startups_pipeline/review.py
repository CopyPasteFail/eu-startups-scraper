from __future__ import annotations

import csv
from pathlib import Path

from openpyxl import Workbook

from .db import Database
from .models import ReviewItem
from .utils import load_json, stable_hash

REVIEW_HEADERS = [
    "review_id",
    "review_type",
    "company",
    "company_url",
    "subject",
    "proposed_value",
    "candidate_values",
    "context",
    "resolution_action",
    "resolution_value",
    "resolution_notes",
]


def make_review_item(
    *,
    review_type: str,
    company_url: str,
    company_name: str,
    subject: str,
    proposed_value: str,
    candidate_values: list[str],
    context: dict[str, str],
) -> ReviewItem:
    review_id = stable_hash(
        review_type, company_url, subject, proposed_value, "|".join(candidate_values)
    )
    return ReviewItem(
        review_id=review_id,
        review_type=review_type,
        company_url=company_url,
        company_name=company_name,
        subject=subject,
        proposed_value=proposed_value,
        candidate_values=candidate_values,
        context=context,
    )


def export_review_files(db: Database, output_dir: Path) -> None:
    rows = db.get_open_review_items()
    csv_path = output_dir / "review_queue.csv"
    xlsx_path = output_dir / "review_queue.xlsx"
    resolutions_path = output_dir / "review_resolutions.csv"
    with csv_path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=REVIEW_HEADERS)
        writer.writeheader()
        for row in rows:
            writer.writerow(
                {
                    "review_id": row["review_id"],
                    "review_type": row["review_type"],
                    "company": row["company_name"],
                    "company_url": row["company_url"],
                    "subject": row["subject"],
                    "proposed_value": row["proposed_value"],
                    "candidate_values": " | ".join(load_json(row["candidate_values_json"], [])),
                    "context": str(load_json(row["context_json"], {})),
                    "resolution_action": "",
                    "resolution_value": "",
                    "resolution_notes": "",
                }
            )
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "review_queue"
    sheet.append(REVIEW_HEADERS)
    for row in rows:
        sheet.append(
            [
                row["review_id"],
                row["review_type"],
                row["company_name"],
                row["company_url"],
                row["subject"],
                row["proposed_value"],
                " | ".join(load_json(row["candidate_values_json"], [])),
                str(load_json(row["context_json"], {})),
                "",
                "",
                "",
            ]
        )
        cell = sheet.cell(row=sheet.max_row, column=4)
        cell.hyperlink = row["company_url"]
        cell.style = "Hyperlink"
    workbook.save(xlsx_path)
    if not resolutions_path.exists():
        with resolutions_path.open("w", newline="", encoding="utf-8") as handle:
            writer = csv.DictWriter(handle, fieldnames=REVIEW_HEADERS)
            writer.writeheader()


def apply_review_resolutions(db: Database, resolutions_path: Path) -> int:
    if not resolutions_path.exists():
        return 0
    applied = 0
    with resolutions_path.open("r", newline="", encoding="utf-8") as handle:
        reader = csv.DictReader(handle)
        for row in reader:
            action = (row.get("resolution_action") or "").strip()
            review_id = (row.get("review_id") or "").strip()
            if not review_id or not action:
                continue
            db.apply_review_resolution(
                review_id=review_id,
                action=action,
                value=(row.get("resolution_value") or "").strip(),
                notes=(row.get("resolution_notes") or "").strip(),
            )
            applied += 1
    return applied
