from __future__ import annotations

import csv
from pathlib import Path

from openpyxl import Workbook

from .db import Database

EXPORT_HEADERS = [
    "Company",
    "EU-link",
    "Company LinkedIn",
    "Person",
    "Person LinkedIn",
    "Role",
    "contacted?",
    "Total funding",
    "Funding stage",
]


def export_results(db: Database, output_dir: Path) -> int:
    rows = db.build_export_rows()
    csv_path = output_dir / "results_clean.csv"
    xlsx_path = output_dir / "results_clean.xlsx"
    with csv_path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(EXPORT_HEADERS)
        for row in rows:
            writer.writerow(
                [
                    row.company,
                    row.eu_link,
                    row.company_linkedin,
                    row.person,
                    row.person_linkedin,
                    row.role,
                    "",
                    row.total_funding,
                    row.funding_stage,
                ]
            )
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "results_clean"
    sheet.append(EXPORT_HEADERS)
    for row in rows:
        sheet.append(
            [
                row.company,
                "EU-link",
                "Company LinkedIn" if row.company_linkedin else "",
                row.person,
                "Person LinkedIn" if row.person and row.person_linkedin else "",
                row.role,
                "",
                row.total_funding,
                row.funding_stage,
            ]
        )
        current = sheet.max_row
        if row.company_website:
            sheet.cell(row=current, column=1).hyperlink = row.company_website
            sheet.cell(row=current, column=1).style = "Hyperlink"
        if row.eu_link:
            sheet.cell(row=current, column=2).hyperlink = row.eu_link
            sheet.cell(row=current, column=2).style = "Hyperlink"
        if row.company_linkedin:
            sheet.cell(row=current, column=3).hyperlink = row.company_linkedin
            sheet.cell(row=current, column=3).style = "Hyperlink"
        if row.person and row.person_linkedin:
            sheet.cell(row=current, column=5).hyperlink = row.person_linkedin
            sheet.cell(row=current, column=5).style = "Hyperlink"
    workbook.save(xlsx_path)
    db.set_metadata("last_export_row_count", str(len(rows)))
    return len(rows)
